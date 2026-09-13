from docx import Document
import json
import sys
from pathlib import Path

def _get_run_info(run) -> dict:
    """Extrai informações de formatação de um run individual."""
    font_name = None
    if run.font and run.font.name:
        font_name = run.font.name
    font_size_pt = None
    if run.font and run.font.size and hasattr(run.font.size, "pt") and run.font.size.pt:
        font_size_pt = run.font.size.pt
    return {
        "text": run.text or "",
        "bold": bool(run.font.bold) if (run.font and run.font.bold is not None) else False,
        "italic": bool(run.font.italic) if (run.font and run.font.italic is not None) else False,
        "font_name": font_name,
        "font_size_pt": font_size_pt,
    }


def _get_paragraph_style_with_runs(para) -> dict:
    """Extrai estilo do parágrafo incluindo lista de runs para formatação inline."""
    runs_info = [_get_run_info(r) for r in (para.runs or []) if r.text]

    if not para.runs:
        font_size = None
        if para.style and hasattr(para.style, "font") and para.style.font and para.style.font.size:
            font_size = para.style.font.size.pt
        return {
            "font_name": "Default",
            "font_size_pt": font_size or 11.0,
            "bold": False,
            "runs": [],
        }

    dominant_run = max(para.runs, key=lambda r: len(r.text or ""))
    font_size = None
    if dominant_run.font and dominant_run.font.size and dominant_run.font.size.pt:
        font_size = dominant_run.font.size.pt
    elif para.style and hasattr(para.style, "font") and para.style.font and para.style.font.size:
        font_size = para.style.font.size.pt

    return {
        "font_name": dominant_run.font.name if (dominant_run.font and dominant_run.font.name) else "Default",
        "font_size_pt": font_size or 11.0,
        "bold": bool(dominant_run.font.bold) if (dominant_run.font and dominant_run.font.bold is not None) else False,
        "runs": runs_info,
    }


_get_paragraph_style = _get_paragraph_style_with_runs

def extract_word_info(docx_path):
    doc = Document(docx_path)
    
    extracted_data = {
        "tables": [],
        "headers": [],
        "footers": []
    }
    
    for table_idx, table in enumerate(doc.tables):
        table_info = {
            "cells": []
        }
        
        for row_idx, row in enumerate(table.rows):
            for col_idx, cell in enumerate(row.cells):
                cell_info = {
                    "row": row_idx,
                    "column": col_idx,
                    "text": cell.text,
                    "paragraphs": []
                }
                
                for para in cell.paragraphs:
                    style = _get_paragraph_style_with_runs(para)
                    para_cell_info = {
                        "text": para.text
                    }
                    para_cell_info.update(style)
                    cell_info["paragraphs"].append(para_cell_info)
                
                table_info["cells"].append(cell_info)
        
        extracted_data["tables"].append(table_info)
    
    for section_idx, section in enumerate(doc.sections):
        if section.header:
            header_info = {
                "paragraphs": [],
                "tables": []
            }
            
            for para in section.header.paragraphs:
                style = _get_paragraph_style_with_runs(para)
                para_info = {"text": para.text}
                para_info.update(style)
                header_info["paragraphs"].append(para_info)
            
            for table in section.header.tables:
                table_data = {
                    "cells": []
                }
                for row_idx, row in enumerate(table.rows):
                    for col_idx, cell in enumerate(row.cells):
                        cell_data = {
                            "row": row_idx,
                            "column": col_idx,
                            "text": cell.text,
                            "paragraphs": []
                        }
                        for para in cell.paragraphs:
                            style = _get_paragraph_style_with_runs(para)
                            para_data = {"text": para.text}
                            para_data.update(style)
                            cell_data["paragraphs"].append(para_data)
                        table_data["cells"].append(cell_data)
                header_info["tables"].append(table_data)
            
            extracted_data["headers"].append(header_info)
        
        if section.footer:
            footer_info = {
                "paragraphs": []
            }
            for para in section.footer.paragraphs:
                style = _get_paragraph_style_with_runs(para)
                para_info = {"text": para.text}
                para_info.update(style)
                footer_info["paragraphs"].append(para_info)
            extracted_data["footers"].append(footer_info)
    
    return extracted_data


def save_to_json(data, output_path):
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


if __name__ == "__main__":
    project_root = Path(__file__).resolve().parent.parent
    docx_file = project_root / "fixtures" / "sample_certificate.docx"
    if not docx_file.exists():
        print("Arquivo 'fixtures/sample_certificate.docx' não encontrado")
        sys.exit(1)
    
    output_json = project_root / "target" / "word_extracted_info.json"
    output_json.parent.mkdir(parents=True, exist_ok=True)
    
    try:
        data = extract_word_info(str(docx_file))
        
        save_to_json(data, output_json)
        print("Processamento concluído")
        
    except Exception as e:
        print(f"[ERRO] Erro ao processar o arquivo: {e}")
        import traceback
        traceback.print_exc()
