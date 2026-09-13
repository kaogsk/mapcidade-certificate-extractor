from docx import Document
from docx.shared import Pt, Mm
import json
import os


def emu_to_mm(emu):
    if emu is None:
        return None
    return round(emu / 914400 * 25.4, 2)


def twips_to_mm(twips):
    if twips is None:
        return None
    return round(twips / 1440 * 25.4, 2)


def extrair_conteudo_celula(celula):
    conteudo = []
    for paragrafo in celula.paragraphs:
        texto = paragrafo.text.strip()
        if texto:
            conteudo.append(texto)
    return "\n".join(conteudo) if conteudo else ""


WNAMESPACE = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def detectar_rowspan(vmerge_val):
    """
    Interpreta o valor de w:vMerge:
    - "restart" → célula inicia rowspan
    - "" (string vazia) → célula continua rowspan
    - None → célula normal, sem rowspan
    """
    if vmerge_val is None:
        return "none"
    if vmerge_val == "restart":
        return "start"
    return "continue"


def _get_vmerge_val(celula):
    """Extrai o valor de w:vMerge de uma célula DOCX. Retorna None se ausente."""
    tc = celula._tc
    tcPr = tc.find(f"{{{WNAMESPACE}}}tcPr")
    if tcPr is None:
        return None
    vmerge = tcPr.find(f"{{{WNAMESPACE}}}vMerge")
    if vmerge is None:
        return None
    val = vmerge.get(f"{{{WNAMESPACE}}}val")
    return val if val is not None else ""


def analisar_tabelas(caminho_arquivo):
    doc = Document(caminho_arquivo)
    resultado = {"tabelas": []}

    for idx_tabela, tabela in enumerate(doc.tables):
        info_tabela = {
            "indice": idx_tabela,
            "num_linhas": len(tabela.rows),
            "num_colunas": len(tabela.columns),
            "colunas": [],
            "linhas": [],
        }

        for idx_col, coluna in enumerate(tabela.columns):
            info_tabela["colunas"].append({
                "indice": idx_col,
                "largura_mm": emu_to_mm(coluna.width),
            })

        rowspan_tracker: dict = {}

        for idx_linha, linha in enumerate(tabela.rows):
            info_linha = {
                "indice": idx_linha,
                "altura_mm": emu_to_mm(linha.height),
                "celulas": [],
            }

            celulas_vistas = set()
            idx_col_logico = 0

            for idx_celula, celula in enumerate(linha.cells):
                celula_id = id(celula._tc)
                if celula_id in celulas_vistas:
                    idx_col_logico += 1
                    continue
                celulas_vistas.add(celula_id)

                colspan = 1
                for idx_prox in range(idx_celula + 1, len(linha.cells)):
                    if id(linha.cells[idx_prox]._tc) == celula_id:
                        colspan += 1
                    else:
                        break

                vmerge_val = _get_vmerge_val(celula)
                rowspan_status = detectar_rowspan(vmerge_val)

                if rowspan_status == "start":
                    rowspan_tracker[idx_col_logico] = {
                        "count": 1,
                        "idx_linha_inicial": idx_linha,
                    }
                    current_rowspan = 1
                elif rowspan_status == "continue":
                    if idx_col_logico in rowspan_tracker:
                        rowspan_tracker[idx_col_logico]["count"] += 1
                        new_count = rowspan_tracker[idx_col_logico]["count"]
                        ini = rowspan_tracker[idx_col_logico]["idx_linha_inicial"]
                        for linha_anterior in resultado["tabelas"][-1]["linhas"] if resultado["tabelas"] else []:
                            for c in linha_anterior["celulas"]:
                                if (c.get("indice_coluna_inicial") == idx_col_logico and
                                        c.get("indice_linha_inicial") == ini):
                                    c["rowspan"] = new_count
                    idx_col_logico += colspan
                    continue
                else:
                    current_rowspan = 1
                    rowspan_tracker.pop(idx_col_logico, None)

                info_celula = {
                    "indice_coluna_inicial": idx_col_logico,
                    "indice_coluna_final": idx_col_logico + colspan - 1,
                    "colspan": colspan,
                    "rowspan": current_rowspan,
                    "indice_linha_inicial": idx_linha,
                    "conteudo": extrair_conteudo_celula(celula),
                    "largura_mm": emu_to_mm(celula.width) if hasattr(celula, "width") else None,
                }

                info_linha["celulas"].append(info_celula)
                idx_col_logico += colspan

            info_tabela["linhas"].append(info_linha)

        resultado["tabelas"].append(info_tabela)

    return resultado


def main():
    projeto_raiz = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    caminho_arquivo = os.path.join(projeto_raiz, "fixtures", "sample_certificate.docx")
    
    print(f"Analisando arquivo: {caminho_arquivo}")
    print("=" * 80)
    
    try:
        resultado = analisar_tabelas(caminho_arquivo)
        
        print(f"\nTotal de tabelas encontradas: {len(resultado['tabelas'])}\n")
        
        for tabela in resultado["tabelas"]:
            print(f"\n{'='*80}")
            print(f"TABELA {tabela['indice']}")
            print(f"{'='*80}")
            print(f"Dimensões: {tabela['num_linhas']} linhas x {tabela['num_colunas']} colunas")
            
            print("\nLarguras das Colunas:")
            for col in tabela["colunas"]:
                print(f"  Coluna {col['indice']}: {col['largura_mm']} mm")
            
            print("\nConteúdo:")
            for linha in tabela["linhas"]:
                print(f"\n  Linha {linha['indice']} (altura: {linha['altura_mm']} mm):")
                for celula in linha["celulas"]:
                    conteudo = celula['conteudo'][:50] + "..." if len(celula['conteudo']) > 50 else celula['conteudo']
                    
                    if celula['colspan'] > 1:
                        info_col = f"Col {celula['indice_coluna_inicial']}-{celula['indice_coluna_final']} (mesclada {celula['colspan']} cols)"
                    else:
                        info_col = f"Col {celula['indice_coluna_inicial']}"
                    
                    print(f"    {info_col}: {conteudo}")
        
        arquivo_saida = os.path.join(projeto_raiz, "target", "tables_extracted.json")
        with open(arquivo_saida, 'w', encoding='utf-8') as f:
            json.dump(resultado, f, ensure_ascii=False, indent=2)
        
        print(f"\n{'='*80}")
        print(f"[OK] Analise completa salva em: {arquivo_saida}")
        
    except FileNotFoundError:
        print(f"ERRO: Arquivo não encontrado: {caminho_arquivo}")
    except Exception as e:
        print(f"ERRO ao processar arquivo: {str(e)}")


if __name__ == "__main__":
    main()
